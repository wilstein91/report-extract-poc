"""손으로 채운 엑셀 정답표를 data/truth/ground_truth.csv 로 옮긴다.

    uv run python scripts/xlsx_to_csv.py
"""

import csv
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "truth" / "ground_truth_작성용.xlsx"
DST = ROOT / "data" / "truth" / "ground_truth.csv"


def main():
    ws = load_workbook(SRC, data_only=True)["정답표"]
    rows = [
        ["" if c is None else str(c).strip() for c in row]
        for row in ws.iter_values() if any(c is not None for c in row)
    ] if hasattr(ws, "iter_values") else [
        ["" if c.value is None else str(c.value).strip() for c in row]
        for row in ws.iter_rows() if any(c.value is not None for c in row)
    ]

    # 엑셀에서 빈칸으로 둔 자리는 빈 문자열로 남긴다.
    # 0 이나 '-' 로 바꾸지 않는다 — '없다'와 '0'은 다른 답이다.
    with DST.open("w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(rows)

    filled = sum(1 for r in rows[1:] if any(r[2:9]))
    print(f"옮겼습니다: {DST}")
    print(f"  전체 {len(rows) - 1}행 중 항목이 하나라도 채워진 행: {filled}")
    blanks = [r[0] for r in rows[1:] if not r[5] and not r[6]]
    if blanks:
        print(f"  투자의견·목표주가가 모두 빈칸인 행: {', '.join(blanks)}")


if __name__ == "__main__":
    main()
